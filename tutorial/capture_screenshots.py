"""
VOLUTE tutorial screen captures
Drives VOLUTE through the steps of the tutorial in docs/tutorial/, on the Nancy
excerpt, and saves the captures its pages show, in English and in French, into
docs/tutorial/images/<language>/.

Each language runs in a process of its own, with SAM2.1 Large. VOLUTE's
settings are read from a copy of volute_config.yaml, so the run changes none of
them, and the optional progress animation is left out, as in a default
installation. The frames are extracted under /Users/Shared/VOLUTE-tutorial on
macOS, a path that names nobody, and deleted afterwards.

The windows open on the screen while it runs, with the platform's own look;
--offscreen draws them off screen instead, with Qt's generic look.

Run from VOLUTE's own environment, where frame extraction works, with the
screen kept awake throughout: a locked screen leaves VOLUTE in the background,
its controls greyed.

    python tutorial/capture_screenshots.py              # both languages
    python tutorial/capture_screenshots.py --lang fr
    python tutorial/capture_screenshots.py --pages start
"""

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
TOOLS = HERE.parent
DOCS = TOOLS / "docs" / "tutorial"
VIDEO = DOCS / "nancy_P1_S1_excerpt.mp4"
LANGUAGES = ["en", "fr"]

# The models VOLUTE is started with: (family, configuration, checkpoint, SAM2++
# task). Paths are relative to the SAM2 installation, as VOLUTE.py gives them
MODELS = {
    "sam2": ("sam2", "configs/sam2.1/sam2.1_hiera_l.yaml",
             "./checkpoints/sam2.1_hiera_large.pt", None),
    "sam2plus-point": ("sam2plus", "configs/sam2.1/sam2.1_hiera_b+_predmasks_decoupled_MAME.yaml",
                       "./checkpoints/SAM2-Plus/checkpoint_phase123.pt", "point"),
}
WINDOW_SIZE = (1280, 800)

# Text the screenshots show beside VOLUTE's own
TEXT = {
    "en": {"object": "Vocal tract", "tongue": "Tongue", "apex": "Tongue apex",
           "palate": "Hard palate", "wall": "Pharyngeal wall", "frame": "Frame {}",
           "time": "Time (s)", "distance": "Tongue – hard palate distance (pixels)"},
    "fr": {"object": "Conduit vocal", "tongue": "Langue", "apex": "Apex lingual",
           "palate": "Palais dur", "wall": "Paroi pharyngale", "frame": "Trame {}",
           "time": "Temps (s)", "distance": "Distance langue – palais dur (pixels)"},
}

# The tutorial's first object, the vocal tract, on the first frame: a box from
# the lips to the lower pharynx, its upper edge below the hard palate, (x0, y0,
# x1, y1) in image pixels
VOCAL_TRACT_BOX = (80, 180, 320, 390)
# Then each group of clicks, in image pixels, is followed by a prediction, the
# first group being empty: the box alone. Positive clicks (left button) mark
# the object, negative ones (right button) what to leave out.
VOCAL_TRACT_STEPS = [
    [],                  # the box alone
    [(185, 188, "+")],   # in the air of the oral cavity
    [(200, 280, "-")],   # on the tongue
    [(300, 300, "+")],   # in the air of the lower pharynx
]
# The second object, the tongue, from a box alone, (x0, y0, x1, y1) in image
# pixels
TONGUE_BOX = (110, 200, 290, 370)
# What the vocal tract of the third page wrongly takes in, marked on frame 1:
# the alveolar region, the soft palate, the nasopharynx behind it
VOCAL_TRACT_PROBLEMS = [(1, 118, 185), (2, 243, 203), (3, 286, 188)]
# The corrections, each followed by a prediction. The second, on the
# nasopharynx, lies above the box, and VOLUTE asks for confirmation
VOCAL_TRACT_CORRECTIONS = [(252, 190, "-"), (290, 175, "-"), (118, 185, "-")]
# The frame-1 export, retouched by hand in GIMP, in each language: its layers
# carry the object names
CORRECTED_XCF = {"en": HERE / "nancy_P1_S1_excerpt_00000_en_CORRECTED.xcf",
                 "fr": HERE / "nancy_P1_S1_excerpt_00000_fr_CORRIGE.xcf"}
# Following both objects through the video: on frame 63 the tongue spills
# into the air of the pharynx, behind its root. The view is zoomed in on that
# region, (x0, y0, x1, y1) in image pixels, a negative point marks the spill,
# and the tongue is re-propagated over a few frames each way
TRACKING_ERROR_FRAME = 62
TRACKING_ZOOM = (200, 170, 350, 320)
TRACKING_CORRECTION = (272, 242)
TRACKING_SPAN = 5
# Two fixed reference points, in image pixels: on the hard palate, and on the
# posterior wall of the pharynx
REFERENCE_POINTS = [("palate", 170, 175), ("wall", 321, 295)]
DISPLAY_FRAME = 74
MONTAGE_FRAMES = [0, 37, 74, 110]
# The tongue apex on frame 1, followed by SAM2++ point tracking; the frames shown
# include 38, one of the two on which the model loses it, and 43 and 78, where the
# tongue draws back. The marker is enlarged, to stand out on the whole image
TONGUE_APEX = (125, 250)
POINT_MONTAGE_FRAMES = [0, 37, 42, 77]
TRACKED_POINT_SIZE = 16
# Frames per second of the excerpt, for the time axis of the distance chart
FRAME_RATE = 50


# How long a capture waits for the window to become active, in seconds: for
# the application to come back to the front by itself, then for a click
REFOCUS_WAIT = 60
INACTIVE_WAIT = 300


def _activate_application():
    """Asks macOS to bring this application to the front, through AppKit's
    activateIgnoringOtherApps:, which Qt's activateWindow() does not call."""
    if sys.platform != "darwin":
        return
    import ctypes
    import ctypes.util
    objc = ctypes.cdll.LoadLibrary(ctypes.util.find_library("objc"))
    objc.objc_getClass.restype = ctypes.c_void_p
    objc.sel_registerName.restype = ctypes.c_void_p
    send = objc.objc_msgSend
    send.restype = ctypes.c_void_p
    send.argtypes = [ctypes.c_void_p, ctypes.c_void_p]
    application = send(objc.objc_getClass(b"NSApplication"),
                       objc.sel_registerName(b"sharedApplication"))
    send.argtypes = [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_bool]
    send(application, objc.sel_registerName(b"activateIgnoringOtherApps:"), True)


def _application_is_active():
    """Whether macOS has this application in front, as AppKit says: the
    controls it draws itself, such as radio buttons, look greyed otherwise."""
    if sys.platform != "darwin":
        return True
    import ctypes
    import ctypes.util
    objc = ctypes.cdll.LoadLibrary(ctypes.util.find_library("objc"))
    objc.objc_getClass.restype = ctypes.c_void_p
    objc.sel_registerName.restype = ctypes.c_void_p
    send = objc.objc_msgSend
    send.restype = ctypes.c_void_p
    send.argtypes = [ctypes.c_void_p, ctypes.c_void_p]
    application = send(objc.objc_getClass(b"NSApplication"),
                       objc.sel_registerName(b"sharedApplication"))
    send.restype = ctypes.c_bool
    return bool(send(application, objc.sel_registerName(b"isActive")))


def _frontmost_application():
    """The name of the application macOS has in front, for the report of a
    capture taken in the background: loginwindow when the screen is locked."""
    if sys.platform != "darwin":
        return ""
    import ctypes
    import ctypes.util
    objc = ctypes.cdll.LoadLibrary(ctypes.util.find_library("objc"))
    ctypes.cdll.LoadLibrary(ctypes.util.find_library("AppKit"))
    objc.objc_getClass.restype = ctypes.c_void_p
    objc.sel_registerName.restype = ctypes.c_void_p
    send = objc.objc_msgSend
    send.restype = ctypes.c_void_p
    send.argtypes = [ctypes.c_void_p, ctypes.c_void_p]
    workspace = send(objc.objc_getClass(b"NSWorkspace"), objc.sel_registerName(b"sharedWorkspace"))
    application = send(workspace, objc.sel_registerName(b"frontmostApplication"))
    if not application:
        return ""
    name = send(application, objc.sel_registerName(b"localizedName"))
    send.restype = ctypes.c_char_p
    return send(name, objc.sel_registerName(b"UTF8String")).decode()


def work_root():
    if sys.platform == "darwin":
        return Path("/Users/Shared/VOLUTE-tutorial")
    return Path(tempfile.gettempdir()) / "VOLUTE-tutorial"


def handoff_file(language, name):
    """Where a session leaves a file for a later one: the project it saved, or
    the positions the point-tracking session exported. Outside work_root(),
    which each session empties."""
    return work_root().parent / f"VOLUTE-tutorial-{language}-{name}"


class Session:
    """One run of VOLUTE in one language, and the helpers that drive it."""

    def __init__(self, language, out_dir, keep_xcf=None, model="sam2"):
        self.language = language
        self.model = model
        self.keep_xcf = keep_xcf
        self.text = TEXT[language]
        self.out = out_dir
        self.out.mkdir(parents=True, exist_ok=True)
        self.work = work_root()
        shutil.rmtree(self.work, ignore_errors=True)
        self.work.mkdir(parents=True)
        shutil.copy(VIDEO, self.work)
        self.video = self.work / VIDEO.name
        self.handlers = []

    # -- Start-up ---------------------------------------------------------

    def start(self):
        import sam2
        sam2_root = Path(sam2.__file__).resolve().parent.parent
        # The SAM2 configuration and checkpoint paths are relative to it
        os.chdir(sam2_root)
        os.environ["PYTORCH_ENABLE_MPS_FALLBACK"] = "1"
        sys.path.insert(0, str(TOOLS))

        config = self.work / "volute_config.yaml"
        text = (TOOLS / "volute_config.yaml").read_text(encoding="utf-8")
        text = text.replace('default_language: "en"', f'default_language: "{self.language}"')
        # Qt's file dialogs can be driven, unlike the system's; none is captured
        text = text.replace("native_file_dialogs: true", "native_file_dialogs: false")
        config.write_text(text, encoding="utf-8")

        import volute.config_manager as config_manager
        import volute.progress_animation as progress_animation

        def load_copy(manager, path=str(config)):
            manager.config_file_path = path
            manager._load_yaml_config(path)

        config_manager.ConfigManager._load_config = load_copy
        progress_animation.find_animation = lambda: None

        from PyQt5.QtCore import QTimer
        from PyQt5.QtWidgets import QApplication
        self.app = QApplication.instance() or QApplication(sys.argv)
        self.app.setApplicationName("VOLUTE")
        self.app.setApplicationDisplayName("VOLUTE")
        # The icon message boxes show, as VOLUTE.py sets it
        from VOLUTE import find_app_icon
        icon = find_app_icon(TOOLS)
        if icon is not None:
            self.app.setWindowIcon(icon)
        self.timer = QTimer()
        self.timer.timeout.connect(self._watch_dialogs)
        self.timer.start(100)

        from volute import SAM2VideoSegmentationApp
        family, config_path, checkpoint, task = MODELS[self.model]
        self.window = SAM2VideoSegmentationApp(
            debug_mode=False, config_path=config_path, checkpoint_path=checkpoint,
            sam2plus_task=task, model_type=family)
        self.window.resize(*WINDOW_SIZE)
        self.window.show()
        self.window.raise_()
        self.ui = self.window.ui_manager
        self.canvas = self.ui.canvas
        self.pump(1.5)

    def close(self):
        import threading
        # Should closing the window hang, the process ends all the same
        threading.Timer(60, lambda: os._exit(0)).start()
        shutil.rmtree(self.work, ignore_errors=True)
        self.window.close()
        self.pump(0.3)

    # -- Event loop and dialogs -------------------------------------------

    def pump(self, seconds=0.3):
        end = time.time() + seconds
        while time.time() < end:
            self.app.processEvents()
            time.sleep(0.01)

    def wait_until(self, condition, timeout=600):
        end = time.time() + timeout
        while not condition():
            if time.time() > end:
                raise TimeoutError("condition not met")
            self.pump(0.1)

    def on_dialog(self, matches, handler):
        """Calls handler(dialog) on the next modal dialog for which matches(dialog)
        holds. A message box no handler claims is closed with its default button."""
        self.handlers.append((matches, handler))

    def _watch_dialogs(self):
        from PyQt5.QtCore import QTimer
        from PyQt5.QtWidgets import QMessageBox
        dialog = self.app.activeModalWidget()
        if dialog is None or getattr(dialog, "_tutorial_seen", False):
            return
        for index, (matches, handler) in enumerate(self.handlers):
            if matches(dialog):
                dialog._tutorial_seen = True
                del self.handlers[index]
                # Once drawn, not from within the timer that found it
                QTimer.singleShot(400, lambda: handler(dialog))
                return
        if isinstance(dialog, QMessageBox):
            dialog._tutorial_seen = True
            QTimer.singleShot(300, dialog.accept)

    # -- Input ------------------------------------------------------------

    def canvas_point(self, x, y):
        """Widget position of image pixel (x, y) on the canvas."""
        from PyQt5.QtCore import QPoint
        dx, dy = self.canvas.axes.transData.transform((x, y))
        ratio = self.canvas.device_pixel_ratio
        return QPoint(int(dx / ratio), int(self.canvas.height() - dy / ratio))

    def click_image(self, x, y, positive=True):
        from PyQt5.QtCore import Qt
        from PyQt5.QtTest import QTest
        button = Qt.LeftButton if positive else Qt.RightButton
        QTest.mouseClick(self.canvas, button, Qt.NoModifier, self.canvas_point(x, y))
        self.pump(0.3)

    def drag_image(self, x0, y0, x1, y1):
        """Left-button drag on the canvas, between image pixels (x0, y0) and
        (x1, y1). The moves are sent as events holding the button, which
        QTest.mouseMove does not."""
        from PyQt5.QtCore import QEvent, Qt
        from PyQt5.QtGui import QMouseEvent
        from PyQt5.QtTest import QTest
        QTest.mousePress(self.canvas, Qt.LeftButton, Qt.NoModifier, self.canvas_point(x0, y0))
        self.pump(0.1)
        for step in range(1, 11):
            point = self.canvas_point(x0 + (x1 - x0) * step / 10, y0 + (y1 - y0) * step / 10)
            event = QMouseEvent(QEvent.MouseMove, point, Qt.NoButton, Qt.LeftButton, Qt.NoModifier)
            self.app.sendEvent(self.canvas, event)
            self.pump(0.03)
        QTest.mouseRelease(self.canvas, Qt.LeftButton, Qt.NoModifier, self.canvas_point(x1, y1))
        self.pump(0.3)

    def click(self, widget):
        from PyQt5.QtCore import Qt
        from PyQt5.QtTest import QTest
        QTest.mouseClick(widget, Qt.LeftButton)
        self.pump(0.3)

    def control(self, name):
        return self.ui.get_control(name)

    def show_tab(self, index):
        self.ui._tab_widget.setCurrentIndex(index)
        self.pump(0.3)

    def go_to_frame(self, index):
        self.control("frame_spinbox").setValue(index + 1)
        self.pump(0.5)

    # -- Screenshots ------------------------------------------------------

    def save(self, widget, name, callouts=None, down_to=None, crop=None, callout_radius=15):
        """Saves widget as name.webp and returns the image.

        callouts lists (number, widget, anchor) triples, each marked with a
        numbered disc of callout_radius points; anchor is (fx, fy, dx, dy), a position given as fractions
        of the widget's size plus an offset in points. down_to keeps the top of
        the widget only, down to the bottom of that child widget; crop is a box
        (left, top, right, bottom) in pixels of the grabbed image."""
        from PIL import Image
        from PyQt5.QtCore import QPoint
        self._present(widget)
        pixmap = widget.grab()
        ratio = pixmap.devicePixelRatio()
        raw = self.work / f"{name}.png"
        pixmap.save(str(raw))
        image = Image.open(raw).convert("RGB")
        if callouts:
            image = self._draw_callouts(image, widget, callouts, ratio, callout_radius)
        if down_to is not None:
            bottom = down_to.mapTo(widget, QPoint(0, down_to.height())).y() + 8
            image = image.crop((0, 0, image.width, min(image.height, int(bottom * ratio))))
        if crop is not None:
            image = image.crop(crop)
        image.save(self.out / f"{name}.webp", quality=90, method=6)
        # The pages show dialogs at their size on screen, the width printed here
        print(f"  {self.language}/{name}.webp  ({round(image.width / ratio)} points wide)")
        return image

    def _present(self, widget):
        """Brings widget's window to the front, active, so that it is drawn
        with the platform's accent colours rather than greyed out, and scrolls
        the control panel back to its top.

        macOS keeps another application in front while it is being used, and
        does not let one bring itself forward: the window is then made active
        for Qt alone, which draws what Qt draws as active, but not the
        controls AppKit draws, such as radio buttons, which the capture
        reports. Should that fail too, the capture waits, up to
        INACTIVE_WAIT seconds, for a click on the window, and goes on greyed
        out after that, saying so."""
        from PyQt5.QtWidgets import QApplication, QScrollArea
        window = widget.window()
        window.raise_()
        window.activateWindow()
        _activate_application()
        for area in self.window.findChildren(QScrollArea):
            area.verticalScrollBar().setValue(0)
        self.pump(0.4)
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen":
            return
        # Another application often comes forward only for a moment
        end = time.time() + REFOCUS_WAIT
        while not _application_is_active() and time.time() < end:
            _activate_application()
            self.pump(1.0)
        if not _application_is_active():
            print(f"  application in the background, behind {_frontmost_application()}: "
                  "its native controls will look greyed", flush=True)
        if window.isActiveWindow():
            return
        QApplication.setActiveWindow(window)
        self.pump(0.4)
        if window.isActiveWindow():
            return
        print("  VOLUTE is not the active window: click on it to go on", flush=True)
        end = time.time() + INACTIVE_WAIT
        while not window.isActiveWindow() and time.time() < end:
            self.pump(0.5)
        if not window.isActiveWindow():
            print("  still inactive: this screenshot will look greyed out", flush=True)
        self.pump(0.4)

    def save_image_area(self, name, marks=None):
        """Saves the canvas without its margins: the image alone. marks lists
        (number, x, y) triples, numbered discs placed at image pixels (x, y)."""
        ratio = self.canvas.device_pixel_ratio
        extent = self.canvas.axes.get_window_extent()
        height = self.canvas.height() * ratio
        box = (int(extent.x0), int(height - extent.y1), int(extent.x1), int(height - extent.y0))
        callouts = None
        if marks:
            callouts = []
            for number, x, y in marks:
                point = self.canvas_point(x, y)
                callouts.append((number, self.canvas, (0, 0, point.x(), point.y())))
        return self.save(self.canvas, name, callouts=callouts, crop=box, callout_radius=10)

    def _draw_callouts(self, image, widget, callouts, ratio, radius_points):
        from PIL import ImageDraw
        from PyQt5.QtCore import QPoint
        draw = ImageDraw.Draw(image)
        radius = int(radius_points * ratio)
        font = self._font(int(radius_points * 1.2 * ratio))
        for number, target, (fx, fy, dx, dy) in callouts:
            point = target.mapTo(widget, QPoint(int(target.width() * fx + dx),
                                                int(target.height() * fy + dy)))
            x, y = point.x() * ratio, point.y() * ratio
            draw.ellipse((x - radius, y - radius, x + radius, y + radius),
                         fill=(126, 80, 173), outline=(255, 255, 255), width=int(2 * ratio))
            draw.text((x, y), str(number), fill=(255, 255, 255), font=font, anchor="mm")
        return image

    @staticmethod
    def _font(size):
        from PIL import ImageFont
        for path in ("/System/Library/Fonts/Helvetica.ttc",
                     "/System/Library/Fonts/SFNS.ttf",
                     "C:/Windows/Fonts/arialbd.ttf",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
            if os.path.exists(path):
                return ImageFont.truetype(path, size)
        return ImageFont.load_default()

    def montage(self, images, labels, name):
        """Places images side by side, each with its label underneath."""
        from PIL import Image, ImageDraw
        tile = 600
        images = [image.resize((tile, round(image.height * tile / image.width)), Image.LANCZOS)
                  for image in images]
        gap, label_height = 16, 44
        width = sum(image.width for image in images) + gap * (len(images) - 1)
        height = max(image.height for image in images) + label_height
        sheet = Image.new("RGB", (width, height), (255, 255, 255))
        draw = ImageDraw.Draw(sheet)
        font = self._font(26)
        x = 0
        for image, label in zip(images, labels):
            sheet.paste(image, (x, 0))
            draw.text((x + image.width / 2, image.height + label_height / 2), label,
                      fill=(40, 40, 40), font=font, anchor="mm")
            x += image.width + gap
        sheet.save(self.out / f"{name}.webp", quality=90, method=6)
        print(f"  {self.language}/{name}.webp")

    # -- Pages ------------------------------------------------------------

    def page_start(self):
        """Before you start: the model chooser and the main window."""
        from volute.model_dialog import ModelSelectionDialog
        chooser = ModelSelectionDialog(self.window.localization, self.window.config_manager,
                                       self.window)
        chooser.show()
        self.pump(0.5)
        self.save(chooser, "start-model-chooser")
        chooser.reject()

        self.save(self.window, "start-main-window", callouts=[
            (1, self.canvas, (0.5, 0.5, 0, 0)),
            (2, self._last_tool_button(), (1, 0.5, 26, 0)),
            (3, self.control("image_slider"), (0.5, 0.5, 0, 0)),
            (4, self.ui._tab_widget.tabBar(), (0, 0.5, -28, 0)),
        ])

    def _last_tool_button(self):
        """The rightmost button of the image toolbar."""
        from PyQt5.QtWidgets import QToolButton
        buttons = [button for button in self.ui.toolbar.findChildren(QToolButton)
                   if button.isVisible()]
        return max(buttons, key=lambda button: button.x())

    def page_load(self):
        """Loading the frames: extraction from the video, then navigation."""
        from PyQt5.QtWidgets import QMessageBox
        from volute.video_frames_ui import VideoFramesDialog

        def fill_and_accept(dialog):
            dialog.video_edit.setText(str(self.video))
            dialog.video_edit.editingFinished.emit()
            self.pump(1.0)
            self.save(dialog, "load-extract-dialog")
            dialog.accept()

        def load_offer(box):
            self.save(box, "load-extract-done")
            box.done(QMessageBox.Yes)

        self.on_dialog(lambda d: isinstance(d, VideoFramesDialog), fill_and_accept)
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.No, load_offer)
        self.window.extract_video_frames()
        self.wait_until(lambda: self.window.image_manager.has_images())
        self.pump(1.0)
        self.save(self.window, "load-frames-loaded")

        self.go_to_frame(60)
        navigation = self.control("image_slider").parentWidget()
        self.save(navigation, "load-navigation")
        self.go_to_frame(0)

    def page_first_object(self):
        """A first object: the vocal tract on the first frame, from a box and a
        few clicks."""
        name_edit = self.control("obj_name_edit")
        name_edit.setText(self.text["object"])
        name_edit.editingFinished.emit()
        self.pump(0.3)
        self.save(self.ui._tab_widget, "object-objects-tab",
                  down_to=self.control("obj_name_edit").parentWidget())

        self.control("define_box_radio").setChecked(True)
        self.pump(0.2)
        self.drag_image(*VOCAL_TRACT_BOX)
        self.save_image_area("object-box")
        self.control("add_points_radio").setChecked(True)

        for step, clicks in enumerate(VOCAL_TRACT_STEPS, 1):
            self.show_tab(self.window.TAB_OBJECTS_POINTS)
            for x, y, sign in clicks:
                self.click_image(x, y, positive=(sign == "+"))
            self.show_tab(self.window.TAB_PREDICTION)
            if step == 1:
                self.save(self.ui._tab_widget, "object-prediction-tab",
                          down_to=self.control("predict_current_btn").parentWidget())
            self.click(self.control("predict_current_btn"))
            self.wait_until(lambda: 0 in self.window.object_manager.object_masks[1])
            self.pump(1.0)
            self.save_image_area(f"object-step-{step}")

    def save_propagation_montage(self, name):
        images = []
        for index in MONTAGE_FRAMES:
            self.go_to_frame(index)
            images.append(self.save_image_area(f"{name}-{index + 1:03d}"))
        for index in MONTAGE_FRAMES:
            (self.out / f"{name}-{index + 1:03d}.webp").unlink()
        self.montage(images, [self.text["frame"].format(i + 1) for i in MONTAGE_FRAMES], name)

    def add_object(self, name):
        self.show_tab(self.window.TAB_OBJECTS_POINTS)
        self.click(self.control("add_object_btn"))
        name_edit = self.control("obj_name_edit")
        name_edit.setText(name)
        name_edit.editingFinished.emit()
        self.pump(0.3)
        return self.window.object_manager.current_object_id

    def predict(self, object_id):
        self.show_tab(self.window.TAB_PREDICTION)
        masks = self.window.object_manager.object_masks
        self.click(self.control("predict_current_btn"))
        self.wait_until(lambda: 0 in masks.get(object_id, {}))
        self.pump(1.0)

    def group_box(self, widget):
        """The group box holding widget."""
        from PyQt5.QtWidgets import QGroupBox
        while widget is not None and not isinstance(widget, QGroupBox):
            widget = widget.parentWidget()
        return widget

    def page_display(self):
        """Display options, then two fixed reference points."""
        self.go_to_frame(DISPLAY_FRAME)
        self.show_tab(self.window.TAB_PREDICTION)
        self.control("show_centroids_checkbox").setChecked(True)
        self.control("show_contours_checkbox").setChecked(True)
        self.pump(1.0)
        contours_group = self.group_box(self.control("show_contours_checkbox"))
        self.save(self.ui._tab_widget, "display-options", down_to=contours_group)
        # The labels, drawn at the centre of each mask, would hide the centroids
        self.control("show_object_labels_checkbox").setChecked(False)
        self.pump(0.5)
        self.save_image_area("display-contours")
        self.control("show_object_labels_checkbox").setChecked(True)
        self.control("show_centroids_checkbox").setChecked(False)
        self.control("show_contours_checkbox").setChecked(False)
        self.pump(0.5)

        self.go_to_frame(0)
        self.show_tab(self.window.TAB_REF_POINTS)
        for key, x, y in REFERENCE_POINTS:
            self.click(self.control("add_ref_point_btn"))
            name_edit = self.control("ref_point_name_edit")
            name_edit.setText(self.text[key])
            name_edit.editingFinished.emit()
            self.pump(0.3)
            self.click_image(x, y)
        self.save(self.ui._tab_widget, "display-reference-points-tab",
                  down_to=self.group_box(self.control("ref_points_list")))
        self.go_to_frame(DISPLAY_FRAME)
        self.save_image_area("display-reference-points")

    def select_object(self, row):
        """Clicks the object at that row of the object list."""
        from PyQt5.QtCore import Qt
        from PyQt5.QtTest import QTest
        self.show_tab(self.window.TAB_OBJECTS_POINTS)
        objects_list = self.control("objects_list")
        rect = objects_list.visualItemRect(objects_list.item(row))
        QTest.mouseClick(objects_list.viewport(), Qt.LeftButton, Qt.NoModifier, rect.center())
        self.pump(0.3)

    def page_objects(self):
        """Several objects: the tongue from a box, then both objects followed
        together."""
        self.go_to_frame(0)
        tongue = self.add_object(self.text["tongue"])
        objects_group = self.control("obj_name_edit").parentWidget()
        self.save(self.ui._tab_widget, "objects-second-object", down_to=objects_group)

        self.control("define_box_radio").setChecked(True)
        self.pump(0.2)
        self.drag_image(*TONGUE_BOX)
        self.save_image_area("objects-tongue-box")
        self.control("add_points_radio").setChecked(True)
        self.predict(tongue)
        self.save_image_area("objects-tongue")

    def page_corrections(self):
        """Correcting masks on the first frame: negative points on what the
        vocal tract wrongly takes in, export for GIMP, then import of the frame
        retouched in GIMP."""
        from PyQt5.QtWidgets import QFileDialog, QMessageBox
        from volute.dialogs import GimpExportDialog
        self.go_to_frame(0)
        self.select_object(0)
        self.save_image_area("corrections-problems", marks=VOCAL_TRACT_PROBLEMS)

        def outside_box(box):
            self.save(box, "corrections-point-outside-box")
            box.done(QMessageBox.Yes)

        # The only message box with a check box: the point outside the box
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.checkBox() is not None,
                       outside_box)
        for step, (x, y, sign) in enumerate(VOCAL_TRACT_CORRECTIONS, 1):
            self.show_tab(self.window.TAB_OBJECTS_POINTS)
            self.click_image(x, y, positive=(sign == "+"))
            self.predict(1)
            if step == 2:
                self.save_image_area("corrections-soft-palate")
        self.save_image_area("corrections-corrected")

        # Named after the frame, as VOLUTE suggests
        frame = Path(self.window.image_manager.get_all_image_paths()[0])
        xcf = self.work / (frame.stem + ".xcf")

        def export_dialog(dialog):
            self.save(dialog, "corrections-gimp-export")
            dialog.accept()

        def save_as(dialog):
            # A dialog created by QFileDialog's static functions refuses accept()
            # from Python: its Save button is clicked instead
            from PyQt5.QtWidgets import QDialogButtonBox
            dialog.selectFile(str(xcf))
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()

        self.on_dialog(lambda d: isinstance(d, GimpExportDialog), export_dialog)
        self.on_dialog(lambda d: isinstance(d, QFileDialog), save_as)
        self.window.export_frame_for_gimp()
        self.wait_until(lambda: xcf.exists() and self.app.activeModalWidget() is None, timeout=180)
        self.pump(1.0)
        if self.keep_xcf:
            self.keep_xcf.mkdir(parents=True, exist_ok=True)
            shutil.copy(xcf, self.keep_xcf / f"{xcf.stem}_{self.language}.xcf")

        # The exported file, as saved from GIMP once retouched
        shutil.copy(CORRECTED_XCF[self.language], xcf)

        def open_file(dialog):
            # selectFile() leaves an Open dialog's button disabled: the name is
            # typed in, as a person would
            from PyQt5.QtTest import QTest
            from PyQt5.QtWidgets import QDialogButtonBox, QLineEdit
            dialog.setDirectory(str(xcf.parent))
            name_edit = dialog.findChild(QLineEdit, "fileNameEdit")
            name_edit.clear()
            QTest.keyClicks(name_edit, xcf.name)
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Open).click()

        def box_warning(box):
            self.save(box, "corrections-import-box-warning")
            box.done(QMessageBox.Yes)

        imported = self.window.object_manager.mask_import_origin
        self.on_dialog(lambda d: isinstance(d, QFileDialog), open_file)
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.Cancel, box_warning)
        self.window.import_mask()
        self.wait_until(lambda: 0 in imported.get(1, set()), timeout=180)
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.pump(1.0)
        self.save_image_area("corrections-imported")

    def page_tracking(self):
        """Following both objects through the video from the corrected first
        frame, then a tracking error further on, corrected on its frame and
        re-propagated over a few frames each way."""
        from volute.dialogs import RepropagationSpanDialog

        def progress(dialog):
            self.wait_until(lambda: dialog.value() >= dialog.maximum() * 0.4)
            self.save(dialog, "tracking-progress")

        self.go_to_frame(0)
        self.show_tab(self.window.TAB_PREDICTION)
        self.on_dialog(lambda d: d.metaObject().className() == "AnimatedProgressDialog", progress)
        self.click(self.control("propagate_btn"))
        frames = len(self.window.image_manager.image_paths)
        masks = self.window.object_manager.object_masks
        self.wait_until(lambda: all(len(masks[i]) == frames for i in (1, 2)))
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.pump(2.0)
        self.save_propagation_montage("tracking-propagated")

        tongue = masks[2]
        self.go_to_frame(TRACKING_ERROR_FRAME)
        self.select_object(1)
        # The labels, drawn at the centre of each mask, would hide the contour
        self.control("show_object_labels_checkbox").setChecked(False)
        self.zoom_to(*TRACKING_ZOOM)
        self.save_image_area("tracking-error")
        before = tongue[TRACKING_ERROR_FRAME]
        self.click_image(*TRACKING_CORRECTION, positive=False)
        self.show_tab(self.window.TAB_PREDICTION)
        self.click(self.control("predict_current_btn"))
        self.wait_until(lambda: tongue[TRACKING_ERROR_FRAME] is not before)
        self.pump(1.0)
        self.save_image_area("tracking-corrected")

        def span_dialog(dialog):
            from PyQt5.QtWidgets import QDialogButtonBox
            dialog.forward_spin.setValue(TRACKING_SPAN)
            dialog.backward_spin.setValue(TRACKING_SPAN)
            # Focus on OK, so that no spin box shows its text cursor
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok).setFocus()
            self.pump(0.3)
            self.save(dialog, "tracking-repropagate-dialog")
            dialog.accept()

        # The last frame of the span in each direction
        ends = (TRACKING_ERROR_FRAME - TRACKING_SPAN, TRACKING_ERROR_FRAME + TRACKING_SPAN)
        before = {index: tongue.get(index) for index in ends}
        self.on_dialog(lambda d: isinstance(d, RepropagationSpanDialog), span_dialog)
        self.click(self.control("repropagate_btn"))
        self.wait_until(lambda: all(tongue.get(i) is not before[i] for i in before))
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.pump(2.0)
        self.reset_zoom()
        self.control("show_object_labels_checkbox").setChecked(True)
        self.go_to_frame(0)

    def zoom_to(self, x0, y0, x1, y1):
        """Shows image pixels x0 to x1 and y0 to y1, as the toolbar's zoom tool
        does, VOLUTE keeping that view across frames."""
        axes = self.canvas.axes
        axes.set_xlim(x0, x1)
        axes.set_ylim(y1, y0)
        self.window.capture_global_zoom_state()
        self.canvas.draw()
        self.pump(0.5)

    def reset_zoom(self):
        """Back to the whole image, as the toolbar's Home button does."""
        self.window.reset_global_zoom()
        self.window.update_display()
        self.pump(0.5)

    def page_export(self):
        """Exporting: the Export menu, the images with their masks, and the
        distances from the masks to the reference points, drawn as a chart."""
        from PyQt5.QtWidgets import QFileDialog, QMenu, QMessageBox
        from volute.pair_selection_dialog import PairSelectionDialog
        self.go_to_frame(0)
        menu = next(widget for widget in self.ui._export_actions[0].associatedWidgets()
                    if isinstance(widget, QMenu))
        self.save_menu(menu, "export-menu")

        images = self.work / "export"
        images.mkdir()
        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.open_file(images))
        self.window.export_masked_images()
        self.wait_until(lambda: self.app.activeModalWidget() is None and
                        len(list(images.rglob("*.jpg"))) == len(self.window.image_manager.image_paths))
        frame = Path(self.window.image_manager.get_all_image_paths()[DISPLAY_FRAME]).stem
        masked = next(images.rglob(f"{frame}*.jpg"))
        self.save_raw(masked, "export-masked-image")

        def pairs_dialog(dialog):
            self.save(dialog, "export-pairs-dialog")
            dialog.accept()

        folder = Path(self.window.image_manager.current_folder).name
        distances = self.work / f"{folder}_closest_mask_points.xlsx"
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.Ok and d.standardButtons() & QMessageBox.Cancel,
                       lambda box: box.done(QMessageBox.Ok))
        self.on_dialog(lambda d: isinstance(d, PairSelectionDialog), pairs_dialog)
        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.save_file_as(distances))
        self.window.export_closest_mask_points()
        self.wait_until(lambda: distances.exists() and self.app.activeModalWidget() is None)
        self.distance_chart(distances, "export-distance-chart")

    def save_menu(self, menu, name):
        """Opens menu over the image and saves it as name.webp."""
        self._present(self.window)
        menu.popup(self.window.mapToGlobal(self.canvas.geometry().center()))
        self.pump(0.8)
        # A menu is grabbed as it stands: bringing its window to the front
        # would close it
        raw = self.work / f"{name}.png"
        pixmap = menu.grab()
        pixmap.save(str(raw))
        menu.close()
        self.save_raw(raw, name, pixmap.devicePixelRatio())

    def save_raw(self, path, name, ratio=1):
        """Saves an image file as name.webp, unchanged; ratio is its pixels
        per point, for the width printed."""
        from PIL import Image
        image = Image.open(path).convert("RGB")
        image.save(self.out / f"{name}.webp", quality=90, method=6)
        print(f"  {self.language}/{name}.webp  ({round(image.width / ratio)} points wide)")

    @staticmethod
    def print_rows(export, keep):
        """Prints the rows of an exported .xlsx file for which keep(row) holds,
        row being a dict by column name, for the page's table."""
        from openpyxl import load_workbook
        rows = list(load_workbook(export, read_only=True).active.values)
        print("  " + " | ".join(rows[0]))
        for values in rows[1:]:
            if keep(dict(zip(rows[0], values))):
                print("  " + " | ".join(str(value) for value in values))

    def distance_chart(self, export, name):
        """Draws the tongue to hard palate distance, frame by frame, from the
        exported file, and prints the pair's first rows for the page's table."""
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from openpyxl import load_workbook
        rows = list(load_workbook(export, read_only=True).active.values)
        header, rows = rows[0], rows[1:]
        column = {key: index for index, key in enumerate(header)}
        pair = sorted((row for row in rows if row[column["Object_Name"]] == self.text["tongue"]
                       and row[column["Target_Name"]] == self.text["palate"]),
                      key=lambda row: row[column["Frame_Index"]])
        print("  " + " | ".join(header))
        for row in pair[:4]:
            print("  " + " | ".join(str(value) for value in row))
        points = [(row[column["Frame_Index"]], row[column["Distance_px"]]) for row in pair]
        figure, axes = plt.subplots(figsize=(8, 3), dpi=150)
        axes.plot([frame / FRAME_RATE for frame, _ in points], [d for _, d in points],
                  color="#7e50ad", linewidth=2)
        axes.set_xlabel(self.text["time"])
        axes.set_ylabel(self.text["distance"])
        axes.set_ylim(bottom=0)
        axes.grid(alpha=0.3)
        for side in ("top", "right"):
            axes.spines[side].set_visible(False)
        figure.tight_layout()
        raw = self.work / f"{name}.png"
        figure.savefig(raw)
        plt.close(figure)
        self.save_raw(raw, name)

    # -- Saving and resuming ------------------------------------------------

    def project_path(self):
        """The project file, named as VOLUTE suggests, after the frame folder."""
        folder = Path(self.window.image_manager.current_folder).name
        return self.work / f"{folder}_objects.volute"

    def page_project_save(self):
        """Saving, first part: the project saved, imported masks included."""
        from PyQt5.QtWidgets import QFileDialog, QMessageBox
        project = self.project_path()

        def imported_masks(box):
            self.save(box, "project-imported-masks")
            box.done(QMessageBox.Yes)

        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.save_file_as(project))
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.No, imported_masks)
        self.window.export_inference_state()
        self.wait_until(lambda: project.exists() and self.app.activeModalWidget() is None)
        shutil.copy(project, handoff_file(self.language, project.name))

    def page_project_resume(self):
        """Saving, second part, in a new VOLUTE: the frames and the project
        opened again, then the masks propagated anew."""
        from PyQt5.QtWidgets import QFileDialog, QMessageBox
        self.load_frames_quietly()
        project = self.project_path()
        shutil.copy(handoff_file(self.language, project.name), project)
        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.open_file(project))
        # Importing replaces the session's objects, once confirmed
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.No, lambda box: box.done(QMessageBox.Yes))
        self.window.import_inference_state()
        objects = self.window.object_manager
        self.wait_until(lambda: len(objects.object_names) >= 2
                        and self.app.activeModalWidget() is None)
        self.pump(1.0)
        self.save_image_area("project-resumed")

        self.show_tab(self.window.TAB_PREDICTION)
        self.click(self.control("propagate_btn"))
        frames = len(self.window.image_manager.image_paths)
        masks = objects.object_masks
        self.wait_until(lambda: all(len(masks.get(i, {})) == frames for i in (1, 2)))
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.pump(2.0)
        self.go_to_frame(DISPLAY_FRAME)
        self.save_image_area("project-propagated")

    def page_batch(self):
        """Saving, third part: the saved project run again by batch
        processing, on the same frames."""
        from volute.batch_dialog import BatchDialog
        output = self.work / "batch"

        def run(dialog):
            dialog._add_row(str(self.project_path()), self.window.image_manager.current_folder, "")
            dialog.output_path_edit.setText(str(output))
            self.pump(0.5)
            self.click(dialog.run_btn)
            self.wait_until(lambda: dialog.worker is not None and dialog.worker.isFinished()
                            and dialog.run_btn.isEnabled(), timeout=1800)
            self.pump(1.0)
            self.save(dialog, "project-batch-dialog")
            for path in sorted(output.rglob("*")):
                if path.is_file() and path.parent.name != "masked_images":
                    print(f"  batch output: {path.relative_to(output)}")
            dialog.reject()

        self.on_dialog(lambda d: isinstance(d, BatchDialog), run)
        self.window.open_batch_dialog()
        self.pump(0.5)

    # -- Settings -----------------------------------------------------------

    def page_settings(self):
        """Settings: three tabs of the settings window, and the Language menu."""
        from PyQt5.QtWidgets import QTabWidget
        from volute.settings_dialog import TABS, SettingsDialog
        dialog = SettingsDialog(self.window.localization, self.window.config_manager,
                                self.window)
        # The capture's own copy of the settings turns the system file
        # dialogs off; the window shows the default instead
        dialog.widgets["ui.native_file_dialogs"].setChecked(True)
        tabs = dialog.findChild(QTabWidget)
        dialog.show()
        for index, (key, _) in enumerate(TABS[:3]):
            tabs.setCurrentIndex(index)
            self.pump(0.5)
            self.save(dialog, f"settings-{key}")
        dialog.reject()

        title = self.window.localization.get_text("language")
        menu = next(action.menu() for action in self.window.menuBar().actions()
                    if action.menu() is not None and action.text() == title)
        self.save_menu(menu, "settings-language-menu")

    # -- Point tracking (a VOLUTE started with SAM2++ in point mode) ---------

    def load_frames_quietly(self):
        """Extracts and loads the frames as page_load does, without its
        screenshots: the frames then bear the same names as in the session
        of the other pages."""
        import contextlib
        import io
        out = self.out
        self.out = self.work / "unused"
        self.out.mkdir(exist_ok=True)
        with contextlib.redirect_stdout(io.StringIO()):
            self.page_load()
        self.out = out

    def save_file_as(self, path):
        """A handler for a Qt Save dialog, saving to path."""
        def handler(dialog):
            # A dialog created by QFileDialog's static functions refuses accept()
            # from Python: its Save button is clicked instead
            from PyQt5.QtWidgets import QDialogButtonBox
            dialog.selectFile(str(path))
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        return handler

    def open_file(self, path):
        """A handler for a Qt Open dialog, opening path."""
        def handler(dialog):
            # selectFile() leaves an Open dialog's button disabled: the name is
            # typed in, as a person would
            from PyQt5.QtTest import QTest
            from PyQt5.QtWidgets import QDialogButtonBox, QLineEdit
            dialog.setDirectory(str(path.parent))
            name_edit = dialog.findChild(QLineEdit, "fileNameEdit")
            name_edit.clear()
            QTest.keyClicks(name_edit, path.name)
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Open).click()
        return handler

    def page_points_tracking(self):
        """Following points, first part: the tongue apex tracked by SAM2++,
        then its positions exported for the session of the other pages."""
        from PyQt5.QtCore import Qt
        from PyQt5.QtWidgets import QFileDialog
        from volute.model_dialog import ModelSelectionDialog
        chooser = ModelSelectionDialog(self.window.localization, self.window.config_manager,
                                       self.window)
        items = chooser.list_widget
        for row in range(items.count()):
            if items.item(row).data(Qt.UserRole) == "sam2plus_point":
                items.setCurrentRow(row)
        chooser.show()
        self.pump(0.5)
        self.save(chooser, "points-model-chooser")
        chooser.reject()

        self.load_frames_quietly()
        name_edit = self.control("obj_name_edit")
        name_edit.setText(self.text["apex"])
        name_edit.editingFinished.emit()
        self.pump(0.3)
        self.click_image(*TONGUE_APEX)
        self.show_tab(self.window.TAB_PREDICTION)
        tracked = self.window.object_manager.object_tracked_points
        self.click(self.control("predict_current_btn"))
        self.wait_until(lambda: 0 in tracked.get(1, {}))
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.control("tracked_point_size_slider").setValue(TRACKED_POINT_SIZE)
        self.pump(1.0)
        # Callout 1 just after the title's first line, which ends with the mode
        title = self.ui.title_label
        metrics = title.fontMetrics()
        half_line = metrics.horizontalAdvance(title.text().split("\n")[0]) / 2
        self.save(self.window, "points-main-window", callouts=[
            (1, title, (0.5, 0.5, half_line + 22, -metrics.lineSpacing() / 2)),
            (2, self.control("predict_current_btn"), (0, 0.5, -20, 0)),
        ])

        self.click(self.control("propagate_btn"))
        self.pump(3.0)
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.pump(2.0)
        images = []
        for index in POINT_MONTAGE_FRAMES:
            self.go_to_frame(index)
            images.append(self.save_image_area(f"points-tracked-{index + 1:03d}"))
            (self.out / f"points-tracked-{index + 1:03d}.webp").unlink()
        self.montage(images, [self.text["frame"].format(i + 1) for i in POINT_MONTAGE_FRAMES],
                     "points-tracked")

        # Named after the frame folder, as VOLUTE suggests
        folder = Path(self.window.image_manager.current_folder).name
        export = self.work / f"{folder}_tracked_points.xlsx"
        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.save_file_as(export))
        self.window.export_tracked_points()
        self.wait_until(lambda: export.exists() and self.app.activeModalWidget() is None)
        shutil.copy(export, handoff_file(self.language, export.name))

    def page_points_import(self):
        """Following points, second part: the apex positions imported beside
        the masks of the other pages."""
        from PyQt5.QtWidgets import QFileDialog, QMessageBox
        folder = Path(self.window.image_manager.current_folder).name
        tracked = self.work / f"{folder}_tracked_points.xlsx"
        shutil.copy(handoff_file(self.language, tracked.name), tracked)

        def count_warning(box):
            self.save(box, "points-import-count-warning")
            box.done(QMessageBox.Yes)

        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.open_file(tracked))
        self.on_dialog(lambda d: isinstance(d, QMessageBox) and d.standardButtons()
                       & QMessageBox.No, count_warning)
        self.window.import_tracked_points()
        self.wait_until(lambda: self.window.imported_point_manager.get_names())
        self.wait_until(lambda: self.app.activeModalWidget() is None)
        self.show_tab(self.window.TAB_PREDICTION)
        self.control("imported_point_marker_size_slider").setValue(TRACKED_POINT_SIZE)
        self.pump(1.0)
        self.save(self.ui._tab_widget, "points-imported-panel",
                  down_to=self.group_box(self.control("imported_point_marker_size_slider")))
        self.go_to_frame(POINT_MONTAGE_FRAMES[2])
        self.save_image_area("points-imported")

        from volute.pair_selection_dialog import PairSelectionDialog
        analysis = self.work / f"{folder}_point_mask_analysis.xlsx"
        self.on_dialog(lambda d: isinstance(d, PairSelectionDialog), lambda d: d.accept())
        self.on_dialog(lambda d: isinstance(d, QFileDialog), self.save_file_as(analysis))
        self.window.export_point_mask_analysis()
        self.wait_until(lambda: analysis.exists() and self.app.activeModalWidget() is None)
        self.print_rows(analysis, lambda row: row["Object_Name"] == self.text["tongue"]
                        and row["Frame_Index"] in POINT_MONTAGE_FRAMES)


PAGES = ["start", "load", "object", "objects", "corrections", "tracking", "display", "export",
         "project", "points", "settings"]


def run_language(language, out_root, pages, keep_xcf=None, offscreen=False):
    """Runs, each in a process of its own, the VOLUTE sessions the pages
    need: the main one, then, from the project page on, a second one that
    reopens the saved project; the point-tracking page adds a SAM2++ one."""
    last = max(PAGES.index(page) for page in pages)

    def run_part(part):
        command = [sys.executable, "-u", __file__, "--lang", language, "--out", str(out_root),
                   "--pages", ",".join(pages), "--part", part]
        if keep_xcf:
            command += ["--keep-xcf", str(keep_xcf)]
        if offscreen:
            command.append("--offscreen")
        subprocess.run(command, check=True)

    run_part("main")
    if last >= PAGES.index("points"):
        run_part("tracking")
    if last >= PAGES.index("project"):
        run_part("resume")


def run_part(part, language, out_root, pages, keep_xcf=None):
    """One VOLUTE session: "main" for the pages up to saving the project,
    "tracking" for SAM2++ point tracking, "resume" for the pages from the
    reopened project on. Each page from the second on builds on all those
    before it."""
    last = max(PAGES.index(page) for page in pages)
    session = Session(language, out_root / language, keep_xcf,
                      model="sam2plus-point" if part == "tracking" else "sam2")
    session.start()
    try:
        if part == "main":
            if "start" in pages:
                session.page_start()
            steps = [("load", session.page_load), ("object", session.page_first_object),
                     ("objects", session.page_objects), ("corrections", session.page_corrections),
                     ("tracking", session.page_tracking), ("display", session.page_display), ("export", session.page_export),
                     ("project", session.page_project_save)]
        elif part == "tracking":
            steps = [("points", session.page_points_tracking)]
        else:
            steps = [("project", session.page_project_resume), ("project", session.page_batch),
                     ("points", session.page_points_import), ("settings", session.page_settings)]
        for page, run in steps:
            if PAGES.index(page) <= last:
                run()
    finally:
        session.close()


def main():
    parser = argparse.ArgumentParser(description="Capture the tutorial's screenshots.")
    parser.add_argument("--lang", choices=LANGUAGES, help="one language only (default: both)")
    parser.add_argument("--out", type=Path, default=DOCS / "images",
                        help="root of the image folders (default: docs/tutorial/images)")
    parser.add_argument("--pages", default=",".join(PAGES),
                        help=f"comma-separated pages to capture (default: {','.join(PAGES)})")
    parser.add_argument("--keep-xcf", type=Path, metavar="FOLDER",
                        help="keep there the .xcf the corrections page exports for GIMP, "
                             "to take GIMP's own screenshots by hand")
    parser.add_argument("--offscreen", action="store_true",
                        help="draw the windows off screen, with Qt's generic look")
    # One of the sessions run_language() starts
    parser.add_argument("--part", choices=["main", "tracking", "resume"], help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.offscreen:
        os.environ["QT_QPA_PLATFORM"] = "offscreen"

    pages = args.pages.split(",")
    unknown = set(pages) - set(PAGES)
    if unknown:
        parser.error(f"unknown pages: {', '.join(sorted(unknown))}")
    keep_xcf = args.keep_xcf.resolve() if args.keep_xcf else None
    if args.part:
        run_part(args.part, args.lang, args.out.resolve(), pages, keep_xcf)
        # Worker threads of Qt or PyTorch can keep the interpreter from exiting
        # once the window is closed; everything is saved by then
        sys.stdout.flush()
        os._exit(0)
    # VOLUTE reads the language at start-up: one language after the other
    for language in [args.lang] if args.lang else LANGUAGES:
        run_language(language, args.out.resolve(), pages, keep_xcf, args.offscreen)


if __name__ == "__main__":
    main()
